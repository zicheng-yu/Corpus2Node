"""Per-format ingest adapters: file -> list of text blocks.

A single registry keyed by extension; each adapter returns text blocks that the
shared chunker turns into EvidenceChunks. Heavy/optional parsers are imported
lazily so a missing one gives a clear message. Document, PDF, image, video and
audio adapters are registered here behind the same extraction interface.
"""
from __future__ import annotations

import csv as _csv
import json as _json
from pathlib import Path

from corpus2node.core.text import normalize_text
from corpus2node.core.types import SourceFile, SourceKind

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}  # Kimi vision: png/jpeg/webp/gif
_VIDEO_EXTS = {".mp4", ".mpeg", ".mpg", ".mov", ".avi", ".flv", ".webm", ".wmv", ".3gp", ".3gpp"}
_AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".aac", ".ogg"}


def extract_blocks_for(source: SourceFile) -> list[str]:
    """Extract text blocks from a source file by its extension."""
    ext = Path(source.filename).suffix.lower()
    adapter = _ADAPTERS.get(ext)
    if adapter is None:
        raise RuntimeError(
            f"Unsupported file type {ext!r}. Supported: {', '.join(sorted(_ADAPTERS))}."
        )
    blocks = [block for block in adapter(source.filename, source.storage_path) if normalize_text(block)]
    if not blocks:
        raise RuntimeError(f"No extractable text found in {source.filename}.")
    return blocks


def kind_for(filename: str) -> SourceKind:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return SourceKind.pdf
    if ext in _AUDIO_EXTS:
        return SourceKind.audio
    if ext in _VIDEO_EXTS:
        return SourceKind.video
    if ext in _IMAGE_EXTS:
        return SourceKind.image
    return SourceKind.document


# ── adapters ────────────────────────────────────────────────────────────────

def _read_text(filename: str, path: str) -> list[str]:
    return [Path(path).read_text(encoding="utf-8", errors="replace")]


def _read_scientific_xml(filename: str, path: str) -> list[str]:
    from corpus2node.scientific.parsers import parse_scientific_xml

    document = parse_scientific_xml(path, source_id=Path(filename).stem)
    blocks = []
    for section in document.sections:
        text = " ".join(value.text for value in section.sentences)
        if text:
            blocks.append(f"{section.title}\n{text}" if section.title else text)
    for table in document.tables:
        rows: dict[int, list[str]] = {}
        for cell in table.cells:
            rows.setdefault(cell.row, []).append(cell.text)
        table_text = "\n".join(" | ".join(cells) for _, cells in sorted(rows.items()))
        if table_text:
            blocks.append(f"{table.caption}\n{table_text}" if table.caption else table_text)
    return blocks


def _read_docx(filename: str, path: str) -> list[str]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("DOCX ingestion needs `python-docx`.") from exc
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ["\n\n".join(parts)]


def _read_pptx(filename: str, path: str) -> list[str]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PPTX ingestion needs `python-pptx`.") from exc
    presentation = Presentation(path)
    blocks: list[str] = []
    for slide in presentation.slides:  # one block per slide (slide = natural unit)
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            blocks.append("\n".join(texts))
    return blocks


def _read_csv(filename: str, path: str) -> list[str]:
    rows: list[str] = []
    with open(path, newline="", encoding="utf-8", errors="replace") as handle:
        for row in _csv.reader(handle):
            cells = [cell.strip() for cell in row if cell.strip()]
            if cells:
                rows.append(" | ".join(cells))
    return ["\n".join(rows)]


def _read_json(filename: str, path: str) -> list[str]:
    data = _json.loads(Path(path).read_text(encoding="utf-8", errors="replace"))
    return [_json.dumps(data, ensure_ascii=False, indent=2)]


def _read_yaml(filename: str, path: str) -> list[str]:
    try:
        import yaml
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("YAML ingestion needs `pyyaml`.") from exc
    data = yaml.safe_load(Path(path).read_text(encoding="utf-8", errors="replace"))
    return [_json.dumps(data, ensure_ascii=False, indent=2) if data is not None else ""]


def _read_pdf(filename: str, path: str) -> list[str]:
    # Moonshot's Files API (file-extract, OCR included) when vision is bound to
    # Kimi; otherwise the fully-local pypdf text layer — offline setups never
    # need a cloud credential for text PDFs.
    from corpus2node.ingest.kimi_client import vision_is_moonshot

    if vision_is_moonshot():
        from corpus2node.ingest.pdf_kimi import extract_pdf_blocks

        return extract_pdf_blocks(filename, path)
    from corpus2node.ingest.pdf_local import extract_pdf_blocks_local

    return extract_pdf_blocks_local(filename, path)


def _read_image(filename: str, path: str) -> list[str]:
    from corpus2node.ingest.image_kimi import describe_image

    return describe_image(filename, path)


def _read_video(filename: str, path: str) -> list[str]:
    from corpus2node.ingest.video_kimi import describe_video

    return describe_video(filename, path)


def _read_audio(filename: str, path: str) -> list[str]:
    from corpus2node.ingest.audio_whisper import transcribe_audio

    return transcribe_audio(filename, path)


_ADAPTERS = {
    ".pdf": _read_pdf,
    ".txt": _read_text,
    ".md": _read_text,
    ".markdown": _read_text,
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    ".csv": _read_csv,
    ".json": _read_json,
    ".yaml": _read_yaml,
    ".yml": _read_yaml,
    ".xml": _read_scientific_xml,
    ".nxml": _read_scientific_xml,
    **{ext: _read_image for ext in _IMAGE_EXTS},
    **{ext: _read_video for ext in _VIDEO_EXTS},
    **{ext: _read_audio for ext in _AUDIO_EXTS},
}

SUPPORTED_EXTENSIONS = frozenset(_ADAPTERS)
