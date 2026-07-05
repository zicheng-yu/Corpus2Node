from __future__ import annotations

import importlib.util
import json
import sys
import types
from pathlib import Path

import pytest

from corpus2node.core.types import SourceFile, SourceKind
from corpus2node.ingest import adapters


def _src(filename: str, path: Path) -> SourceFile:
    return SourceFile(
        kind=SourceKind.document, filename=filename, content_type="x",
        storage_path=str(path), size_bytes=path.stat().st_size,
    )


def test_markdown_adapter(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# 标题\n\n二叉搜索树 用于 高效 查找。", encoding="utf-8")
    blocks = adapters.extract_blocks_for(_src("a.md", p))
    assert blocks and "二叉搜索树" in blocks[0]


def test_docx_adapter(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("二叉搜索树是一种树结构。")
    document.add_paragraph("用于高效查找。")
    p = tmp_path / "a.docx"
    document.save(str(p))
    blocks = adapters.extract_blocks_for(_src("a.docx", p))
    assert blocks and "二叉搜索树" in blocks[0]


def test_pptx_adapter(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "树结构 用于 组织 层次 数据"
    p = tmp_path / "a.pptx"
    presentation.save(str(p))
    blocks = adapters.extract_blocks_for(_src("a.pptx", p))
    assert blocks and any("树结构" in block for block in blocks)


def test_csv_adapter(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("名称,定义\n二叉搜索树,用于查找的树\n", encoding="utf-8")
    blocks = adapters.extract_blocks_for(_src("a.csv", p))
    assert blocks and "二叉搜索树" in blocks[0]


def test_json_adapter(tmp_path):
    p = tmp_path / "a.json"
    p.write_text(json.dumps({"concept": "二叉搜索树"}, ensure_ascii=False), encoding="utf-8")
    blocks = adapters.extract_blocks_for(_src("a.json", p))
    assert blocks and "二叉搜索树" in blocks[0]


def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "a.exe"
    p.write_bytes(b"\x00\x01")
    with pytest.raises(RuntimeError):
        adapters.extract_blocks_for(_src("a.exe", p))


def test_kind_for_maps_extensions():
    assert adapters.kind_for("a.pdf") == SourceKind.pdf
    assert adapters.kind_for("a.md") == SourceKind.document
    assert adapters.kind_for("a.docx") == SourceKind.document
    assert adapters.kind_for("a.png") == SourceKind.image
    assert adapters.kind_for("a.mp3") == SourceKind.audio
    assert adapters.kind_for("a.mp4") == SourceKind.video


def test_multimodal_extensions_are_registered():
    assert ".png" in adapters.SUPPORTED_EXTENSIONS
    assert ".mp3" in adapters.SUPPORTED_EXTENSIONS
    assert ".mp4" in adapters.SUPPORTED_EXTENSIONS


def _seed_vision_kimi() -> None:
    """Bind a Kimi-style credential to the `vision` purpose in the isolated registry."""
    from corpus2node.llm import store
    from corpus2node.llm.credentials import (
        LLMSettings,
        ProviderCredential,
        ProviderKind,
        Purpose,
        PurposeBinding,
    )

    cred = ProviderCredential(
        credential_id="kimi1", label="Kimi", kind=ProviderKind.openai,
        base_url="https://api.moonshot.cn/v1", api_key="test-key", default_model="kimi-k2.6",
    )
    store.save(LLMSettings(credentials=[cred], bindings={Purpose.vision: PurposeBinding(credential_id="kimi1")}))


def test_image_adapter_requires_vision_binding(tmp_path):
    # No `vision` binding in the (isolated, empty) registry → clear error, no network.
    p = tmp_path / "a.png"
    p.write_bytes(b"\x89PNG\r\n")
    with pytest.raises(RuntimeError, match="vision"):
        adapters.extract_blocks_for(_src("a.png", p))


def test_audio_dependency_is_bundled():
    # Audio is transcribed locally (the Kimi API has no audio input), so faster-whisper
    # is a core dependency now — audio ingestion must work without any opt-in extra.
    assert importlib.util.find_spec("faster_whisper") is not None
    assert ".mp3" in adapters.SUPPORTED_EXTENSIONS


def test_audio_error_is_clear_if_transcriber_missing(tmp_path, monkeypatch):
    # If the local transcriber is somehow unavailable, the error must point at the real
    # fix (it's local; Kimi can't do audio), not a stale `--extra audio` instruction.
    monkeypatch.setitem(sys.modules, "faster_whisper", None)  # force ImportError on import
    p = tmp_path / "a.mp3"
    p.write_bytes(b"\x00\x01")
    with pytest.raises(RuntimeError, match="locally"):
        adapters.extract_blocks_for(_src("a.mp3", p))


def _tiny_pdf(text: str) -> bytes:
    """Assemble a minimal one-page PDF with computed xref offsets (valid, no deps)."""
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    return bytes(out)


def test_pdf_falls_back_to_local_extraction_without_moonshot(tmp_path):
    # vision unbound (isolated registry) → the fully-offline pypdf path, no network.
    p = tmp_path / "a.pdf"
    p.write_bytes(_tiny_pdf("Binary search trees enable fast lookup"))
    blocks = adapters.extract_blocks_for(_src("a.pdf", p))
    assert blocks and "Binary search trees" in blocks[0]


def test_pdf_uses_kimi_when_vision_is_moonshot(tmp_path, monkeypatch):
    import corpus2node.ingest.pdf_kimi as pdf_kimi

    _seed_vision_kimi()
    monkeypatch.setattr(pdf_kimi, "extract_pdf_blocks", lambda filename, path: ["kimi 提取内容"])
    p = tmp_path / "a.pdf"
    p.write_bytes(_tiny_pdf("ignored"))
    assert adapters.extract_blocks_for(_src("a.pdf", p)) == ["kimi 提取内容"]


def test_scanned_pdf_error_points_at_kimi(tmp_path):
    from corpus2node.ingest.pdf_local import extract_pdf_blocks_local

    p = tmp_path / "scan.pdf"
    p.write_bytes(_tiny_pdf(""))  # a page with no extractable text ≈ scanned
    with pytest.raises(RuntimeError, match="Kimi"):
        extract_pdf_blocks_local("scan.pdf", p)


def test_video_falls_back_to_local_transcription_without_moonshot(tmp_path, monkeypatch):
    # vision bound to a local provider (no Moonshot Files API) → transcribe the audio track.
    import corpus2node.ingest.audio_whisper as audio_whisper
    from corpus2node.ingest.video_kimi import describe_video
    from corpus2node.llm import store
    from corpus2node.llm.credentials import (
        LLMSettings,
        ProviderCredential,
        ProviderKind,
        Purpose,
        PurposeBinding,
    )

    cred = ProviderCredential(
        credential_id="o1", label="ollama", kind=ProviderKind.ollama, default_model="gemma3:4b"
    )
    store.save(LLMSettings(credentials=[cred], bindings={Purpose.vision: PurposeBinding(credential_id="o1")}))
    monkeypatch.setattr(audio_whisper, "transcribe_audio", lambda filename, path: ["本地转写文本"])

    p = tmp_path / "a.mp4"
    p.write_bytes(b"fake video bytes")
    assert describe_video("a.mp4", str(p)) == ["本地转写文本"]


def test_video_adapter_uses_kimi_file_upload(tmp_path, monkeypatch):
    from corpus2node.ingest.video_kimi import describe_video

    _seed_vision_kimi()  # Kimi credential bound to the vision purpose (was settings.kimi_*)
    p = tmp_path / "a.mp4"
    p.write_bytes(b"fake video bytes")

    calls: dict[str, object] = {}

    class FakeFiles:
        def create(self, *, file, purpose):
            calls["purpose"] = purpose
            calls["uploaded_bytes"] = file.read()
            return types.SimpleNamespace(id="file-video-123")

        def delete(self, file_id):
            calls["deleted"] = file_id

    class FakeCompletions:
        def create(self, **kwargs):
            calls["chat_kwargs"] = kwargs
            return types.SimpleNamespace(
                choices=[types.SimpleNamespace(message=types.SimpleNamespace(content="视频摘要"))]
            )

    class FakeChat:
        def __init__(self):
            self.completions = FakeCompletions()

    class FakeOpenAI:
        def __init__(self, **kwargs):
            calls["client_kwargs"] = kwargs
            self.files = FakeFiles()
            self.chat = FakeChat()

    fake_openai = types.ModuleType("openai")
    fake_openai.OpenAI = FakeOpenAI
    monkeypatch.setitem(sys.modules, "openai", fake_openai)

    blocks = describe_video("a.mp4", str(p))

    assert blocks == ["视频摘要"]
    assert calls["purpose"] == "video"
    assert calls["uploaded_bytes"] == b"fake video bytes"
    assert calls["deleted"] == "file-video-123"
    content = calls["chat_kwargs"]["messages"][1]["content"]
    assert content[0]["video_url"]["url"] == "ms://file-video-123"
